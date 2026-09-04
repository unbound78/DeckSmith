# -*- coding: utf-8 -*-
"""DeckSmith PPT 字体字号合规校验。

交付任何 python-pptx 生成的 PPT 前跑一遍:确保全部文字是统一中文字体且 ≥ 下限字号。
用 ASCII 标记 [OK]/[BAD] 兼容 Windows GBK 终端。

用法:
    python verify_pptx.py <file.pptx> [--font 微软雅黑] [--min 12]
"""
import argparse
from pptx import Presentation
from pptx.util import Pt


def verify(path, font="微软雅黑", min_pt=12):
    prs = Presentation(path)
    bad_font, bad_size, total = [], [], 0
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if not run.text.strip():
                        continue
                    total += 1
                    if run.font.name != font:
                        bad_font.append((i, run.font.name, run.text[:30]))
                    sz = run.font.size
                    if sz is None or sz < Pt(min_pt):
                        shown = f"{sz.pt:.0f}pt" if sz else "未设置"
                        bad_size.append((i, shown, run.text[:30]))

    print(f"\n========== PPT 字体字号校验 ({total} runs) ==========")
    if bad_font:
        print(f"[BAD] 非「{font}」字体 {len(bad_font)} 处:")
        for pg, fn, tx in bad_font:
            print(f"   P{pg}: 字体={fn} | {tx}")
    if bad_size:
        print(f"[BAD] 字号 < {min_pt}pt 或未设置 {len(bad_size)} 处:")
        for pg, sz, tx in bad_size:
            print(f"   P{pg}: 字号={sz} | {tx}")
    if not bad_font and not bad_size:
        print(f"[OK]  字体字号 100% 合规(全部「{font}」+ ≥{min_pt}pt)")
    return len(bad_font) + len(bad_size)


def main():
    ap = argparse.ArgumentParser(description="DeckSmith PPT 字体字号校验")
    ap.add_argument("pptx", help="待校验的 .pptx 路径")
    ap.add_argument("--font", default="微软雅黑", help="要求的字体(默认微软雅黑)")
    ap.add_argument("--min", type=int, default=12, help="字号下限 pt(默认 12)")
    args = ap.parse_args()
    n = verify(args.pptx, args.font, args.min)
    raise SystemExit(1 if n else 0)


if __name__ == "__main__":
    main()
