# -*- coding: utf-8 -*-
"""
DeckSmith · pptx_toolkit
========================
基于 python-pptx + lxml 的 PPT 视觉工具层。提供 python-pptx 原生不直接支持的
能力,全部输出 **PPT 原生可编辑形状**(不是图片):

- 线性 / 径向渐变填充(形状 + 文字渐变)
- 外阴影
- 字体统一(默认微软雅黑 + ≥12pt 下限,解决中文 PPT 字体三套问题)
- 基础形状 + 渐变形状(矩形 / 圆角矩形 / 椭圆 / 梯形)
- 多样式混排文本框
- 文本溢出预检(中文按字宽估算高度,生成前就能发现溢出)

设计原则:
- **纯工具层,零业务内容**——颜色、文案、布局全部由调用方传入
- 通过 lxml 直接操作 DrawingML XML,实现 python-pptx API 缺失的渐变/文字渐变/阴影
- 中文场景友好:set_run_font 统一写入 ea/latin/cs 三套 typeface,避免字体回退

依赖:python-pptx>=0.6.21, lxml>=4.9.0
"""
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

NS_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'

# 默认值(可在调用处覆盖)
DEFAULT_FONT = "微软雅黑"   # 中文场景默认;英文场景可传 font="Calibri" 等
MIN_PT = 12                 # 字号下限,低于此值会被抬到此值(可读性铁律)
DEFAULT_TEXT = RGBColor(0x1A, 0x2A, 0x40)  # 中性深色,仅作 add_text 兜底


# =========================================================
# 色值工具
# =========================================================
def _to_rgb(c):
    """'RRGGBB' 字符串 / (r,g,b) 元组 → (r,g,b) 元组。"""
    if isinstance(c, tuple):
        return c
    s = str(c)
    return (int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))


def _rgb_to_tuple(c):
    if isinstance(c, tuple):
        return c
    return (c[0], c[1], c[2])


# =========================================================
# 渐变 / 阴影(XML 直接操作)
# =========================================================
def _build_gradFill_xml(stops, angle_deg=90, path_circle=False):
    """构建 gradFill XML。stops: [(pos_0_100, RGBColor_or_tuple, alpha_0_1)]。"""
    gs_xml = ''
    for pos, rgb, alpha in stops:
        if hasattr(rgb, '__iter__') and not isinstance(rgb, str):
            r, g, b = rgb[0], rgb[1], rgb[2]
        else:
            r, g, b = _to_rgb(rgb)
        hex_color = f'{r:02X}{g:02X}{b:02X}'
        if alpha < 1.0:
            gs_xml += (f'<a:gs pos="{int(pos*1000)}">'
                       f'<a:srgbClr val="{hex_color}">'
                       f'<a:alpha val="{int(alpha*100000)}"/>'
                       f'</a:srgbClr></a:gs>')
        else:
            gs_xml += f'<a:gs pos="{int(pos*1000)}"><a:srgbClr val="{hex_color}"/></a:gs>'
    if path_circle:
        path_xml = '<a:path path="circle"><a:fillToRect l="50000" t="50000" r="50000" b="50000"/></a:path>'
    else:
        path_xml = f'<a:lin ang="{int(angle_deg*60000)}" scaled="0"/>'
    return (f'<a:gradFill xmlns:a="{NS_A}" flip="none" rotWithShape="1">'
            f'<a:gsLst>{gs_xml}</a:gsLst>{path_xml}</a:gradFill>')


def _find_spPr(shape):
    for child in shape._element:
        if child.tag.endswith('spPr'):
            return child
    return None


def apply_grad_shape(shape, stops, angle_deg=90, path_circle=False):
    """对形状应用渐变填充。stops: [(pos_0_100, RGBColor_or_tuple, alpha_0_1)]。"""
    spPr = _find_spPr(shape)
    if spPr is None:
        return
    for tag in ('solidFill', 'gradFill', 'noFill', 'blipFill', 'pattFill'):
        for el in spPr.findall(qn(f'a:{tag}')):
            spPr.remove(el)
    stops_norm = [(pos, _rgb_to_tuple(c), a) for (pos, c, a) in stops]
    gradFill = etree.fromstring(_build_gradFill_xml(stops_norm, angle_deg, path_circle))
    ln = spPr.find(qn('a:ln'))
    if ln is not None:
        ln.addprevious(gradFill)
    else:
        for tag in ('prstGeom', 'custGeom'):
            geom = spPr.find(qn(f'a:{tag}'))
            if geom is not None:
                geom.addnext(gradFill)
                return
        spPr.append(gradFill)


def apply_grad_run(run, stops, angle_deg=0):
    """对 text run 应用渐变填充(文字渐变)。"""
    rPr = run._r.get_or_add_rPr()
    for tag in ('solidFill', 'gradFill', 'noFill'):
        for el in rPr.findall(qn(f'a:{tag}')):
            rPr.remove(el)
    stops_norm = [(pos, _rgb_to_tuple(c), a) for (pos, c, a) in stops]
    gradFill = etree.fromstring(_build_gradFill_xml(stops_norm, angle_deg))
    ln = rPr.find(qn('a:ln'))
    if ln is not None:
        ln.addnext(gradFill)
    else:
        rPr.insert(0, gradFill)


def apply_shadow(shape, blur_pt=8, dist_pt=2, dir_deg=45, alpha=0.3):
    """对形状应用外阴影。blur/dist 单位 pt,dir 角度(度)。"""
    spPr = _find_spPr(shape)
    if spPr is None:
        return
    for el in spPr.findall(qn('a:effectLst')):
        spPr.remove(el)
    blur_emu = int(blur_pt * 12700)
    dist_emu = int(dist_pt * 12700)
    dir_60000 = int(dir_deg * 60000)
    alpha_v = int(alpha * 100000)
    effect_xml = (f'<a:effectLst xmlns:a="{NS_A}">'
                  f'<a:outerShdw blurRad="{blur_emu}" dist="{dist_emu}" dir="{dir_60000}" algn="tl" rotWithShape="0">'
                  f'<a:srgbClr val="000000"><a:alpha val="{alpha_v}"/></a:srgbClr>'
                  f'</a:outerShdw></a:effectLst>')
    effect = etree.fromstring(effect_xml)
    spPr.append(effect)


# =========================================================
# 字体 / 文本
# =========================================================
def set_run_font(run, size_pt=14, color=None, bold=False, italic=False,
                 font=DEFAULT_FONT, min_pt=MIN_PT):
    """统一设置 run 字体。写入 ea/latin/cs 三套 typeface,避免中文字体回退。"""
    size_pt = max(size_pt, min_pt)
    run.font.name = font
    run.font.size = Pt(size_pt)
    if color is not None:
        run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:latin', 'a:cs'):
        for el in rPr.findall(qn(tag)):
            rPr.remove(el)
    for tag in ('a:ea', 'a:latin', 'a:cs'):
        el = etree.SubElement(rPr, qn(tag))
        el.set('typeface', font)


def add_text(slide, left, top, width, height, content, size=14, color=None,
             bold=False, italic=False, align='left', anchor='top', line_spacing=1.3,
             font=DEFAULT_FONT):
    """添加文本框。content 可为 str,或 [(text, {style}), ...] 实现一行内多样式混排。

    style 支持键:size / color / bold / italic。
    """
    if color is None:
        color = DEFAULT_TEXT
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    if anchor == 'middle':
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    elif anchor == 'bottom':
        tf.vertical_anchor = MSO_ANCHOR.BOTTOM
    para = tf.paragraphs[0]
    if align == 'center':
        para.alignment = PP_ALIGN.CENTER
    elif align == 'right':
        para.alignment = PP_ALIGN.RIGHT
    para.line_spacing = line_spacing
    if isinstance(content, str):
        run = para.add_run()
        run.text = content
        set_run_font(run, size, color, bold, italic, font)
    else:
        for text, st in content:
            run = para.add_run()
            run.text = text
            set_run_font(run, st.get('size', size), st.get('color', color),
                         st.get('bold', bold), st.get('italic', italic), font)
    return box


# =========================================================
# 形状
# =========================================================
def _set_fill_line(shape, fill=None, line=None, line_width=0.75, no_line=False):
    if fill is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    if no_line:
        shape.line.fill.background()
    elif line is not None:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_width)


def add_rect(slide, left, top, width, height, fill=None, line=None, line_width=0.75, no_line=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                   Inches(left), Inches(top), Inches(width), Inches(height))
    _set_fill_line(shape, fill, line, line_width, no_line)
    return shape


def add_round_rect(slide, left, top, width, height, fill=None, line=None,
                   line_width=0.75, corner=0.08, no_line=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   Inches(left), Inches(top), Inches(width), Inches(height))
    shape.adjustments[0] = corner
    _set_fill_line(shape, fill, line, line_width, no_line)
    return shape


def add_oval(slide, left, top, width, height, fill=None, line=None, line_width=1, no_line=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                                   Inches(left), Inches(top), Inches(width), Inches(height))
    _set_fill_line(shape, fill, line, line_width, no_line)
    return shape


def add_trapezoid(slide, left, top, width, height, fill=None, rotation=0, no_line=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.TRAPEZOID,
                                   Inches(left), Inches(top), Inches(width), Inches(height))
    if rotation:
        shape.rotation = rotation
    _set_fill_line(shape, fill, no_line=no_line)
    return shape


def add_grad_rect(slide, left, top, width, height, stops, angle=90, line=None, line_width=0.75, no_line=True):
    shape = add_rect(slide, left, top, width, height, fill=RGBColor(0xFF, 0xFF, 0xFF),
                     line=line, line_width=line_width, no_line=no_line)
    apply_grad_shape(shape, stops, angle_deg=angle)
    return shape


def add_grad_round_rect(slide, left, top, width, height, stops, angle=90,
                        line=None, line_width=0.75, corner=0.08, no_line=True):
    shape = add_round_rect(slide, left, top, width, height, fill=RGBColor(0xFF, 0xFF, 0xFF),
                           line=line, line_width=line_width, corner=corner, no_line=no_line)
    apply_grad_shape(shape, stops, angle_deg=angle)
    return shape


def add_grad_oval(slide, left, top, width, height, stops, angle=135, line=None, no_line=True, path_circle=False):
    shape = add_oval(slide, left, top, width, height, fill=RGBColor(0xFF, 0xFF, 0xFF),
                     line=line, no_line=no_line)
    apply_grad_shape(shape, stops, angle_deg=angle, path_circle=path_circle)
    return shape


def add_grad_trapezoid(slide, left, top, width, height, stops, angle=90, rotation=0):
    shape = add_trapezoid(slide, left, top, width, height, fill=RGBColor(0xFF, 0xFF, 0xFF), rotation=rotation)
    apply_grad_shape(shape, stops, angle_deg=angle)
    return shape


def add_image(slide, path, left, top, width, height):
    return slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(width), Inches(height))


# =========================================================
# 文本溢出预检(生成前发现溢出,中文场景刚需)
# =========================================================
def estimate_text_height(text, w_in, size_pt=12, line_spacing=1.35):
    """估算文字在指定宽度文本框内需要的最小高度(英寸)。

    中文按 0.18 in/字 @ 12pt 保守估计(真实渲染含字间距 + 标点占位);
    line_spacing 是行距倍数;保留上下 0.06 in padding(含底部最后一行 descent)。
    """
    if not text:
        return 0
    char_w = (size_pt / 12.0) * 0.18
    margin_lr = 0.10
    avail_w = max(w_in - margin_lr, 0.1)
    chars_per_line = max(int(avail_w / char_w), 1)
    n_chars = len(text)
    n_lines = (n_chars + chars_per_line - 1) // chars_per_line
    line_h = (size_pt / 72.0) * line_spacing
    return n_lines * line_h + 0.06


def verify_text_fit(label, text, w_in, h_in, size_pt=12, line_spacing=1.35, tol=0.02):
    """校验文字是否能装入文本框。返回 (ok, need_h, msg)。

    使用 ASCII 标记 [OK]/[BAD] 兼容 Windows GBK 终端。
    """
    need_h = estimate_text_height(text, w_in, size_pt, line_spacing)
    if need_h > h_in + tol:
        msg = (f"[BAD] OVERFLOW [{label}]: {len(text)}字 需 {need_h:.2f}in, "
               f"框 {h_in:.2f}in (短 {need_h - h_in:.2f}in)")
        return (False, need_h, msg)
    return (True, need_h, f"[OK]  [{label}]: {len(text)}字 需 {need_h:.2f}in, 框 {h_in:.2f}in")


def audit_text_boxes(cases):
    """批量审计文本框溢出。cases: [(label, text, w_in, h_in, size_pt, line_spacing), ...]。
    返回溢出项数;打印逐项结果 + 修复建议。
    """
    print("\n========== 文本溢出审计 ==========")
    fails = []
    for case in cases:
        label, text, w, h = case[0], case[1], case[2], case[3]
        sz = case[4] if len(case) > 4 else 12
        ls = case[5] if len(case) > 5 else 1.35
        ok, need_h, msg = verify_text_fit(label, text, w, h, sz, ls)
        print(msg)
        if not ok:
            fails.append((label, need_h, h))
    print(f"========== 共 {len(cases)} 项, 溢出 {len(fails)} 项 ==========")
    if fails:
        print("!! 需修复:")
        for label, need, have in fails:
            print(f"   [{label}] 加大本框至 >= {need:.2f}in (差 {need - have:.2f}in), 或缩文字, 或拉长底邻框")
    return len(fails)
